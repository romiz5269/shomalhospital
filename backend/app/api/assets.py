import os
import uuid
from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from openpyxl import load_workbook
from sqlalchemy.orm import Session, selectinload

from app.core.database import get_db
from app.core.security import get_current_user
from app.models import User, Asset, AssetType, AssetHardware, OSInfo, SoftwareApp, AssetImage, PMVisit, AssetUpgrade
from app.schemas import AssetCreate, AssetUpdate, AssetOut
from app.schemas.it import AssetUpgradeCreate, AssetUpgradeOut
from app.services.alerts import compute_asset_alert, sync_all_reminders

router = APIRouter(prefix="/assets", tags=["تجهیزات"])

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "uploads", "asset_images")
os.makedirs(UPLOAD_DIR, exist_ok=True)

# Excel column mapping (old simple format)
IMPORT_COLUMNS_SIMPLE = {
    "نام": "name", "name": "name",
    "نوع": "asset_type", "type": "asset_type",
    "سریال": "serial_number", "serial": "serial_number",
    "ip": "ip_address", "IP": "ip_address",
    "mac": "mac_address",
    "محل": "location", "location": "location",
    "مسئول": "assigned_to",
    "برند": "brand", "مدل": "model",
    "تأمین‌کننده": "vendor_name", "vendor": "vendor_name",
}

# Full Excel column mapping matching user's HAMKAF format
IMPORT_COLUMNS_FULL = {
    "ROW": "_row",
    "MB": "hw_mb",
    "CPU": "hw_cpu",
    "RAM": "hw_ram",
    "VGA": "hw_vga",
    "POWER": "hw_power",
    "HARD": "hw_hard",
    "IP": "ip_address",
    "MailUser": "sw_mail_user",
    "MAC": "mac_address",
    "Asset Code": "asset_code",
    "AntiVirus/Status": "sw_antivirus",
    "OS": "sw_os_name",
    "USER NAME ID": "user_name_id",
    "USER NAME - Used": "assigned_to",
    "SYSTEM NAME": "name",
    "SYSTEM NAME-OLD": "system_name_old",
    "PM Date": "pm_date",
    "Monitor/Asset Code": "hw_monitor",
    "PRINTER/Asset Code": "hw_printer",
    "بخش - Section": "section_name",
    "Section ID": "section_id",
    "طبقه": "floor",
    "عکس مشخصات": "_photo",
}


def _q_assets(db: Session):
    return db.query(Asset).options(
        selectinload(Asset.hardware),
        selectinload(Asset.os_info),
        selectinload(Asset.sw_apps),
        selectinload(Asset.images),
    )


def _latest_visit(db: Session, asset_id: int) -> PMVisit | None:
    return (
        db.query(PMVisit)
        .filter(PMVisit.device_id == asset_id)
        .order_by(PMVisit.visit_date.desc())
        .first()
    )


def _asset_out(asset: Asset, db: Session | None = None, viewer: User | None = None) -> AssetOut:
    status, days = compute_asset_alert(asset)
    out = AssetOut.model_validate(asset)
    out.alert_status = status
    out.days_until_activation = days
    if viewer is None or viewer.role != "admin":
        out.windows_key = None
    if db:
        visit = _latest_visit(db, asset.id)
        if visit:
            out.last_received_date = visit.visit_date
            out.last_return_date = visit.return_date
            out.next_service_date = visit.next_pm_date
    return out


def _upsert_hardware(db: Session, asset: Asset, hw_data: dict | None):
    if not hw_data:
        return
    if asset.hardware:
        for k, v in hw_data.items():
            if v is not None:
                setattr(asset.hardware, k, v)
    else:
        asset.hardware = AssetHardware(device_id=asset.id, **hw_data)
        db.add(asset.hardware)


def _upsert_software(db: Session, asset: Asset, sw_data: dict | None):
    if not sw_data:
        return
    os_payload = {
        "os_name": sw_data.get("os_name"),
        "os_version": sw_data.get("os_version"),
        "windows_key": sw_data.get("windows_key"),
    }
    if asset.os_info:
        for k, v in os_payload.items():
            if v is not None:
                setattr(asset.os_info, k, v)
    else:
        asset.os_info = OSInfo(device_id=asset.id, **os_payload)
        db.add(asset.os_info)

    # Keep software apps normalized
    antivirus = (sw_data.get("antivirus") or "").strip()
    antivirus_status = (sw_data.get("antivirus_status") or "").strip()
    mail_user = (sw_data.get("mail_user") or "").strip()

    if antivirus or antivirus_status:
        av = (
            db.query(SoftwareApp)
            .filter(SoftwareApp.device_id == asset.id, SoftwareApp.app_name == "AntiVirus")
            .first()
        )
        if not av:
            av = SoftwareApp(device_id=asset.id, app_name="AntiVirus")
            db.add(av)
        if antivirus:
            av.version = antivirus
        if antivirus_status:
            av.status = "active" if antivirus_status.upper() == "YES" else "inactive"

    if mail_user:
        mail = (
            db.query(SoftwareApp)
            .filter(SoftwareApp.device_id == asset.id, SoftwareApp.app_name == "Mail")
            .first()
        )
        if not mail:
            mail = SoftwareApp(device_id=asset.id, app_name="Mail")
            db.add(mail)
        mail.version = mail_user
        mail.status = "active"


@router.get("/", response_model=list[AssetOut])
def list_assets(
    department_id: int | None = None,
    asset_type: str | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    q = _q_assets(db)
    if department_id:
        q = q.filter(Asset.section_id == department_id)
    if asset_type:
        q = q.filter(Asset.asset_type == AssetType(asset_type))
    return [_asset_out(a, db, current_user) for a in q.order_by(Asset.name).all()]


@router.post("/import/preview")
async def import_preview(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    content = await file.read()
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)

    preview = []
    errors = []
    detected_format: str | None = None
    any_sheet_had_data = False

    # روی همه‌ی شیت‌های فایل حلقه بزن، نه فقط شیت فعال
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue  # شیت خالیه، ردش کن

        headers = [str(h or "").strip() for h in rows[0]]

        # Detect format: full (HAMKAF) or simple
        is_full = "SYSTEM NAME" in headers or "MB" in headers
        col_map = IMPORT_COLUMNS_FULL if is_full else IMPORT_COLUMNS_SIMPLE

        mapping: dict[int, str] = {}
        for i, h in enumerate(headers):
            key = col_map.get(h) or col_map.get(h.lower())
            if key:
                mapping[i] = key

        if "name" not in mapping.values():
            # این شیت ستون نام قابل‌شناسایی نداره؛ به‌جای خطا دادن، ردش کن و برو شیت بعدی
            continue

        any_sheet_had_data = True
        if detected_format is None:
            detected_format = "full" if is_full else "simple"

        for idx, row in enumerate(rows[1:], start=2):
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue
            item: dict = {}
            for col_i, field in mapping.items():
                val = row[col_i] if col_i < len(row) else None
                item[field] = str(val).strip() if val is not None else ""
            if not item.get("name"):
                errors.append(f"شیت «{ws.title}» - ردیف {idx}: نام خالی")
                continue
            item["_format"] = "full" if is_full else "simple"
            item["_sheet"] = ws.title
            preview.append(item)

    if not any_sheet_had_data:
        raise HTTPException(
            status_code=400,
            detail="ستون «نام»/«SYSTEM NAME» در هیچکدام از شیت‌های فایل پیدا نشد",
        )

    return {
        "preview": preview[:500],
        "total": len(preview),
        "errors": errors,
        "format": detected_format or "simple",
    }


@router.post("/import/confirm")
def import_confirm(
    payload: dict,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    rows = payload.get("rows") or []
    created = 0
    for item in rows:
        name = item.get("name")
        if not name:
            continue
        is_full = item.get("_format") == "full"

        at = item.get("asset_type", "pc").lower() if not is_full else "pc"
        try:
            asset_type = AssetType(at)
        except ValueError:
            asset_type = AssetType.PC

        asset = Asset(
            name=name,
            asset_type=asset_type,
            serial_number=item.get("serial_number") or None,
            ip_address=item.get("ip_address") or None,
            mac_address=item.get("mac_address") or None,
            asset_code=item.get("asset_code") or None,
            assigned_to=item.get("assigned_to") or None,
            user_name_id=item.get("user_name_id") or None,
            system_name_old=item.get("system_name_old") or None,
            section_name=item.get("section_name") or None,
            section_id=item.get("section_id") or None,
            floor=item.get("floor") or None,
            pm_date=item.get("pm_date") or None,
            location=item.get("location") or None,
            brand=item.get("brand") or None,
            model=item.get("model") or None,
            vendor_name=item.get("vendor_name") or None,
        )
        if asset_type in (AssetType.PC, AssetType.SERVER):
            asset.windows_activated_at = datetime.utcnow()
        db.add(asset)
        db.flush()

        if is_full:
            hw = AssetHardware(
                device_id=asset.id,
                mb=item.get("hw_mb") or None,
                cpu=item.get("hw_cpu") or None,
                ram=item.get("hw_ram") or None,
                vga=item.get("hw_vga") or None,
                power=item.get("hw_power") or None,
                hard=item.get("hw_hard") or None,
            )
            monitor_raw = item.get("hw_monitor", "")
            if monitor_raw and monitor_raw != "-":
                hw.monitor_asset_code = monitor_raw
            printer_raw = item.get("hw_printer", "")
            if printer_raw and printer_raw != "-":
                parts = printer_raw.split()
                hw.printer_name = " ".join(parts[:-1]) if len(parts) > 1 else printer_raw
                hw.printer_asset_code = parts[-1] if len(parts) > 1 else None
            db.add(hw)

            os_rec = OSInfo(
                device_id=asset.id,
                os_name=item.get("sw_os_name") or None,
            )
            db.add(os_rec)

            antivirus_raw = item.get("sw_antivirus", "")
            if antivirus_raw:
                av_app = SoftwareApp(
                    device_id=asset.id,
                    app_name="AntiVirus",
                    version=antivirus_raw,
                    status="active" if antivirus_raw.upper() == "YES" else "inactive",
                )
                db.add(av_app)

            mail_user = item.get("sw_mail_user")
            if mail_user:
                mail_app = SoftwareApp(
                    device_id=asset.id,
                    app_name="Mail",
                    version=mail_user,
                    status="active",
                )
                db.add(mail_app)

        created += 1
    db.commit()
    sync_all_reminders(db)
    return {"created": created}


@router.get("/{asset_id}", response_model=AssetOut)
def get_asset(asset_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    asset = _q_assets(db).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="تجهیز یافت نشد")
    return _asset_out(asset, db, current_user)


@router.post("/", response_model=AssetOut, status_code=201)
def create_asset(data: AssetCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    hw_data = data.hardware.model_dump(exclude_unset=True) if data.hardware else None
    sw_data = data.software.model_dump(exclude_unset=True) if data.software else None
    payload = data.model_dump(exclude={"hardware", "software"})
    if "department_id" in payload:
        payload.pop("department_id")
    payload["asset_type"] = AssetType(payload["asset_type"])
    if current_user.role != "admin":
        payload["windows_key"] = None
    asset = Asset(**payload)
    if not asset.windows_activated_at and asset.asset_type in (AssetType.PC, AssetType.SERVER):
        asset.windows_activated_at = datetime.utcnow()
    db.add(asset)
    db.flush()
    if hw_data:
        db.add(AssetHardware(device_id=asset.id, **hw_data))
    if sw_data:
        db.add(OSInfo(device_id=asset.id, os_name=sw_data.get("os_name"), os_version=sw_data.get("os_version"), windows_key=sw_data.get("windows_key")))
    db.commit()
    db.refresh(asset)
    sync_all_reminders(db)
    return _asset_out(_q_assets(db).filter(Asset.id == asset.id).first(), db, current_user)


@router.put("/{asset_id}", response_model=AssetOut)
def update_asset(asset_id: int, data: AssetUpdate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    asset = _q_assets(db).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="تجهیز یافت نشد")
    hw_data = data.hardware.model_dump(exclude_unset=True) if data.hardware else None
    sw_data = data.software.model_dump(exclude_unset=True) if data.software else None
    payload = data.model_dump(exclude_unset=True, exclude={"hardware", "software"})
    payload.pop("department_id", None)
    if current_user.role != "admin":
        payload.pop("windows_key", None)
    for k, v in payload.items():
        if k == "asset_type" and v:
            v = AssetType(v)
        setattr(asset, k, v)
    asset.updated_at = datetime.utcnow()
    _upsert_hardware(db, asset, hw_data)
    _upsert_software(db, asset, sw_data)
    db.commit()
    db.refresh(asset)
    sync_all_reminders(db)
    return _asset_out(asset, db, current_user)


@router.delete("/{asset_id}", status_code=204)
def delete_asset(asset_id: int, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="تجهیز یافت نشد")

    # حذف رکوردهای وابسته قبل از حذف خود asset (جلوگیری از ForeignKeyViolation)
    db.query(SoftwareApp).filter(SoftwareApp.device_id == asset_id).delete()
    db.query(AssetHardware).filter(AssetHardware.device_id == asset_id).delete()
    db.query(OSInfo).filter(OSInfo.device_id == asset_id).delete()
    db.query(AssetUpgrade).filter(AssetUpgrade.device_id == asset_id).delete()

    # عکس‌ها: هم فایل روی دیسک هم رکورد دیتابیس پاک بشه
    images = db.query(AssetImage).filter(AssetImage.device_id == asset_id).all()
    for img in images:
        path = os.path.join(UPLOAD_DIR, img.stored_name)
        if os.path.exists(path):
            os.remove(path)
        db.delete(img)

    db.delete(asset)
    db.commit()


# ── Asset images ──

@router.post("/{asset_id}/images")
async def upload_image(
    asset_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="تجهیز یافت نشد")
    ext = os.path.splitext(file.filename or "img.jpg")[1] or ".jpg"
    stored = f"{uuid.uuid4().hex}{ext}"
    path = os.path.join(UPLOAD_DIR, stored)
    data = await file.read()
    with open(path, "wb") as f:
        f.write(data)
    img = AssetImage(
        device_id=asset_id,
        stored_name=stored,
        original_name=file.filename or "image.jpg",
        content_type=file.content_type or "image/jpeg",
        size_bytes=len(data),
    )
    db.add(img)
    db.commit()
    db.refresh(img)
    return {
        "id": img.id,
        "original_name": img.original_name,
        "stored_name": img.stored_name,
        "content_type": img.content_type,
        "size_bytes": img.size_bytes,
        "created_at": str(img.created_at),
    }


@router.delete("/{asset_id}/images/{image_id}", status_code=204)
def delete_image(asset_id: int, image_id: int, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    img = db.query(AssetImage).filter(AssetImage.id == image_id, AssetImage.device_id == asset_id).first()
    if not img:
        raise HTTPException(status_code=404, detail="عکس یافت نشد")
    path = os.path.join(UPLOAD_DIR, img.stored_name)
    if os.path.exists(path):
        os.remove(path)
    db.delete(img)
    db.commit()


# ── Upgrades ──

@router.get("/{asset_id}/upgrades", response_model=list[AssetUpgradeOut])
def list_upgrades(asset_id: int, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="تجهیز یافت نشد")
    return (
        db.query(AssetUpgrade)
        .filter(AssetUpgrade.device_id == asset_id)
        .order_by(AssetUpgrade.upgraded_at.desc())
        .all()
    )


@router.post("/{asset_id}/upgrades", response_model=AssetUpgradeOut, status_code=201)
def create_upgrade(
    asset_id: int,
    data: AssetUpgradeCreate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="تجهیز یافت نشد")
    upgrade = AssetUpgrade(
        device_id=asset_id,
        title=data.title,
        component=data.component,
        description=data.description,
        upgraded_at=data.upgraded_at or datetime.utcnow(),
        upgraded_by=user.full_name or user.username,
    )
    db.add(upgrade)
    db.commit()
    db.refresh(upgrade)
    return upgrade


@router.delete("/{asset_id}/upgrades/{upgrade_id}", status_code=204)
def delete_upgrade(
    asset_id: int,
    upgrade_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    upgrade = (
        db.query(AssetUpgrade)
        .filter(AssetUpgrade.id == upgrade_id, AssetUpgrade.device_id == asset_id)
        .first()
    )
    if not upgrade:
        raise HTTPException(status_code=404, detail="رکورد آپگرید یافت نشد")
    db.delete(upgrade)
    db.commit()