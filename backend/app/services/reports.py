import socket
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, selectinload

from app.models import Asset, AssetHardware, OSInfo, SoftwareApp, Reminder, PMVisit, InventoryItem

_LTR = Alignment(horizontal="left", vertical="center", wrap_text=True)
_RTL = Alignment(horizontal="right", vertical="center", wrap_text=True)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

_THIN = Side(style="thin", color="B0BEC5")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill(start_color="E8F0FA", end_color="E8F0FA", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="003B8E", size=13)
_BODY_FONT = Font(size=11, color="1A2D45")


def _style_sheet(ws, ltr_columns: set[int] | None = None) -> None:
    ws.sheet_view.rightToLeft = True
    if ws.max_row >= 1:
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 30

    ltr_cols = ltr_columns or set()
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            is_header = cell.row == 1
            cell.border = _BORDER
            cell.font = _HEADER_FONT if is_header else _BODY_FONT
            cell.alignment = _LTR if cell.column in ltr_cols else _RTL
            if is_header:
                cell.fill = _HEADER_FILL

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 24

    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 5, 16), 64)


def get_local_ip() -> str:
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return "127.0.0.1"


def get_network_ips() -> list[str]:
    ips = set()
    try:
        hostname = socket.gethostname()
        for info in socket.getaddrinfo(hostname, None, socket.AF_INET):
            ip = info[4][0]
            if not ip.startswith("127."):
                ips.add(ip)
    except Exception:
        pass
    local = get_local_ip()
    if local != "127.0.0.1":
        ips.add(local)
    return sorted(ips)


def _filter_pm_visits(db: Session, date_from=None, date_to=None, work_type=None):
    q = db.query(PMVisit)
    if date_from:
        q = q.filter(PMVisit.visit_date >= date_from)
    if date_to:
        q = q.filter(PMVisit.visit_date <= date_to)
    if work_type and work_type != "all":
        q = q.filter(PMVisit.work_type == work_type)
    return q.order_by(PMVisit.visit_date.desc()).all()


def _fmt_date(dt) -> str:
    if not dt:
        return ""
    return str(dt)[:10]


def _get_app_ver(db: Session, device_id: int, app_name: str) -> str:
    app = db.query(SoftwareApp).filter(SoftwareApp.device_id == device_id, SoftwareApp.app_name == app_name).first()
    return app.version or "" if app else ""


def _hw_map(db: Session) -> dict[int, AssetHardware]:
    return {h.device_id: h for h in db.query(AssetHardware).all()}


def _sw_map(db: Session) -> dict[int, OSInfo]:
    return {s.device_id: s for s in db.query(OSInfo).all()}


def export_xlsx(db: Session, report_type: str, **filters) -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)

    if report_type in ("assets", "all"):
        hw = _hw_map(db)
        sw = _sw_map(db)
        ws = wb.create_sheet("تجهیزات")
        ws.append([
            "ROW", "MB", "CPU", "RAM", "VGA", "POWER", "HARD",
            "IP", "MailUser", "MAC", "Asset Code",
            "AntiVirus/Status", "OS",
            "USER NAME ID", "USER NAME - Used",
            "SYSTEM NAME", "SYSTEM NAME-OLD",
            "PM Date",
            "Monitor/Asset Code", "PRINTER/Asset Code",
            "بخش - Section", "Section ID", "طبقه",
            "عکس مشخصات",
        ])
        ltr = {1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20}
        row_num = 0
        for a in db.query(Asset).order_by(Asset.id).all():
            row_num += 1
            h = hw.get(a.id)
            s = sw.get(a.id)
            av = ""
            av_apps = db.query(SoftwareApp).filter(SoftwareApp.device_id == a.id, SoftwareApp.app_name == "AntiVirus").first()
            if av_apps:
                av = av_apps.status or av_apps.version or ""
            monitor_info = ""
            if h and (h.monitor_name or h.monitor_asset_code):
                monitor_info = f"{h.monitor_name or ''} {h.monitor_asset_code or ''}".strip()
            elif h and h.monitor_asset_code:
                monitor_info = h.monitor_asset_code
            printer_info = ""
            if h and (h.printer_name or h.printer_asset_code):
                printer_info = f"{h.printer_name or ''} {h.printer_asset_code or ''}".strip()
            ws.append([
                row_num,
                h.mb if h else (a.cpu or ""),
                h.cpu if h else (a.cpu or ""),
                h.ram if h else (a.ram or ""),
                h.vga if h else (a.gpu or ""),
                h.power if h else (a.power_info or ""),
                h.hard if h else (a.storage or ""),
                a.ip_address or "",
                _get_app_ver(db, a.id, "Mail"),
                a.mac_address or "",
                a.asset_code or "",
                av,
                s.os_name if s else (a.os_name or ""),
                a.user_name_id or "",
                a.assigned_to or "",
                a.name,
                a.system_name_old or "",
                a.pm_date or "",
                monitor_info,
                printer_info,
                a.section_name or "",
                a.section_id or "",
                a.floor or "",
                "دارد" if a.images else "",
            ])
        _style_sheet(ws, ltr_columns=ltr)

    if report_type in ("pm_visits", "work_cases", "all"):
        assets_map = {a.id: a.name for a in db.query(Asset).all()}
        ws = wb.create_sheet("عیب‌یابی و تعمیرات")
        ws.append(["نوع", "تاریخ دریافت", "تاریخ تحویل", "تاریخ بازنگری", "تحویل‌گیرنده", "واحد", "سرویس‌دهنده", "تجهیز", "یادداشت"])
        for v in _filter_pm_visits(db, filters.get("date_from"), filters.get("date_to"), filters.get("work_type")):
            wt = getattr(v, "work_type", "pm") or "pm"
            ws.append([
                "دوره‌ای" if wt == "pm" else "موردی",
                _fmt_date(v.visit_date),
                _fmt_date(v.return_date),
                _fmt_date(v.next_pm_date),
                v.recipient_name or "",
                v.recipient_unit or "",
                v.performed_by or "",
                assets_map.get(v.device_id, "") if v.device_id else "",
                v.notes or "",
            ])
        _style_sheet(ws, ltr_columns={2, 3, 4})

    if report_type in ("inventory", "all"):
        ws = wb.create_sheet("کالای جدید")
        ws.append(["نام", "دسته", "سریال", "تعداد", "تاریخ دریافت", "تاریخ نصب", "وضعیت", "کاربرد", "یادداشت"])
        for i in db.query(InventoryItem).order_by(InventoryItem.received_date.desc()).all():
            ws.append([
                i.name,
                i.category,
                i.serial_number or "",
                i.quantity,
                _fmt_date(i.received_date),
                _fmt_date(i.installed_date),
                i.status,
                i.purpose or "",
                i.notes or "",
            ])
        _style_sheet(ws, ltr_columns={5, 6})

    if report_type in ("reminders", "all"):
        ws = wb.create_sheet("هشدارها")
        ws.append(["عنوان", "نوع", "وضعیت", "موعد", "توضیحات"])
        for r in db.query(Reminder).filter(Reminder.is_resolved == False).all():
            ws.append([r.title, r.reminder_type.value, r.status.value, _fmt_date(r.due_date), r.description or ""])
        _style_sheet(ws, ltr_columns={4})

    if not wb.sheetnames:
        ws = wb.create_sheet("گزارش")
        ws.append(["داده‌ای یافت نشد"])
        _style_sheet(ws)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
